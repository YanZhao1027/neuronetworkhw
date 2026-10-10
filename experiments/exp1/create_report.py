#!/usr/bin/env python3
"""Create a visually styled Word report directly from genuine metrics/artifacts."""
import json, math, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
MET=json.loads((ROOT/"results"/"metrics.json").read_text(encoding="utf-8"))
REG=MET.get("regression",{})
MN=MET.get("mnist",{})
FIG=ROOT/"figures"
BASE=ROOT/"results"
DOC=ROOT/"report-template.docx"

def set_font(style,name="Noto Sans CJK SC",size=10.3,bold=False,color=None):
    style.font.name=name;style.font.size=Pt(size);style.font.bold=bold
    if color:style.font.color.rgb=RGBColor(*color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"),name)

doc=Document()
sec=doc.sections[0]
sec.page_height,sec.page_width=Inches(11.69),Inches(8.27)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.86);sec.right_margin=Inches(.86)
set_font(doc.styles["Normal"],size=10.2,color=(41,48,60))
doc.styles["Normal"].paragraph_format.space_after=Pt(6)
doc.styles["Normal"].paragraph_format.line_spacing=1.25
for lev,pt,col in [(1,15.5,(27,64,91)),(2,12.5,(32,83,119)),(3,11,(45,83,103))]:
    st=doc.styles[f"Heading {lev}"];set_font(st,size=pt,bold=True,color=col)
    st.paragraph_format.space_before=Pt(13 if lev==1 else 9)
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True
code_style=doc.styles.add_style("CodeSmall",WD_STYLE_TYPE.PARAGRAPH)
set_font(code_style,"DejaVu Sans Mono",7.7,color=(31,44,58))
code_style.paragraph_format.space_after=Pt(0)
code_style.paragraph_format.line_spacing=1.1
capstyle=doc.styles.add_style("FigureCaption",WD_STYLE_TYPE.PARAGRAPH)
set_font(capstyle,size=9.0,color=(100,108,119))
capstyle.paragraph_format.space_after=Pt(10)

def line(s="",style=None):
    return doc.add_paragraph(s,style=style)
def h(s,lev=1):return doc.add_heading(s,level=lev)
def text(s):return line(s)
def fig(name,caption,width=6.05):
    path=FIG/name
    if not path.exists():
        text("【图像尚未生成】"+name);return
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(path),width=Inches(width))
    c=doc.add_paragraph(caption,style="FigureCaption")
    c.alignment=WD_ALIGN_PARAGRAPH.CENTER
def tbl(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers))
    t.style="Light Shading Accent 1"
    for i,x in enumerate(headers):
        t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for row in t.rows:
        for cell in row.cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for c in t.rows[0].cells:
        tcPr=c._tc.get_or_add_tcPr()
        sh=OxmlElement("w:shd");sh.set(qn("w:fill"),"DDE8F3");tcPr.append(sh)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
def code(s):
    for ss in s.strip("\n").splitlines():
        p=doc.add_paragraph(style="CodeSmall")
        p.add_run(ss)
def page():doc.add_page_break()

p=doc.add_paragraph()
p.paragraph_format.space_before=Pt(46)
r=p.add_run("深度学习实验一")
r.bold=True;r.font.size=Pt(27);r.font.color.rgb=RGBColor(25,67,105)
p=doc.add_paragraph("NumPy 操作 · 正弦回归 · Softmax 手写数字识别")
p.style="Subtitle"
text("2026 · 基础实验与独立拓展研究")
text("课程：深度学习 / 神经网络入门")
tbl(["个人信息（提交前填写）","内容"],[
    ("点名册序号","________"),("班级（与点名册一致）","____________________"),
    ("12 位学号","____________________"),("姓名","____________________"),
    ("实验平台","Ubuntu · NVIDIA RTX 5090 D v2 · PyTorch 2.7.0 + CUDA 12.8")])
text("说明：此报告中的曲线、参数和指标均由本项目代码实际运行产生；拓展内容与作业必做部分明确区分。")
h("一、实验目标与运行环境")
text("本次实验掌握 NumPy 数组与索引、基于梯度下降的线性/多项式回归、MNIST 的 784→10 Softmax 线性分类。实验采用固定随机种子 2026，保留源代码、原始数据、训练日志与结果文件以便复现。")
tbl(["环境项目","实测配置"],[
    ("Python",MET.get("environment",{}).get("python","见程序环境记录")),
    ("PyTorch",MET.get("environment",{}).get("torch","见日志")),
    ("CUDA Runtime",str(MET.get("environment",{}).get("cuda_build","—"))),
    ("计算设备",str(MET.get("environment",{}).get("gpu","CPU"))),
    ("随机种子",str(MET.get("seed","2026")))])
h("二、NumPy 基本操作")
text("操作步骤：导入 NumPy；建立一维与二维数组，检查 type / shape / 下标；生成全零、全一和单位矩阵；使用 arange 建立 12 元素数组，reshape 至 3×4 并执行行、列与负索引切片。以下是代码和真实输出的并排证据。")
fig("01_numpy_code_and_output.png","图 1  NumPy 基础操作：实际源代码与终端输出",width=6.3)
text("重点结果：a.shape=(3,)，b.shape=(2,3)，reshape 后 f.shape=(3,4)；f[1,:]=[4,5,6,7]，f[:,2:] 为最后两列，f[2,-1]=11。所有完整可运行代码见 numpy_exercise.py。")
h("三、回归任务：数据生成与模型训练")
r=REG
if r:
    h("3.1 数据生成",2)
    text("100 个 x 在 [-3,3] 等距采样。目标 y=sin(x)+ε，ε 服从 [-0.5,0.5] 均匀分布。为保证两种模型公平比较，采用完全相同的带噪声样本，数据保存在 results/sine_training_data.csv。")
    fig("02_sine_scatter.png","图 2  正弦信号的 100 个带噪训练样本")
    h("3.2 线性回归",2)
    text("模型：ŷ=wx+b；损失函数：MSE=Σ(ŷ−y)²/n。使用 PyTorch 的 nn.Linear(1,1)、MSELoss 和 Adam 优化器；依次执行清空梯度、前向传播、计算损失、反向传播和参数更新。")
    lin=r["linear"];cub=r["cubic"]
    tbl(["指标","实测结果"],[("w",f"{lin['weights'][0]:.6f}"),
      ("b",f"{lin['bias']:.6f}"),("训练 MSE",f"{lin['mse']:.6f}"),
      ("参数更新次数",lin["steps"])])
    fig("03_linear_regression.png","图 3  线性回归拟合曲线与训练样本")
    text("分析：直线只能刻画输入与输出之间的一阶趋势，无法适应正弦曲线的弯曲部分，属于模型表达能力限制导致的欠拟合；持续增加训练轮次无法根本消除这一结构性误差。")
    h("3.3 三次多项式线性回归",2)
    text("特征映射 φ(x)=[x,x²,x³]，模型 ŷ=w₁x+w₂x²+w₃x³+b。模型对参数仍为线性的，但扩展后的输入表达能力更强。")
    tbl(["指标","实测结果"],[
      ("w₁",f"{cub['weights'][0]:.6f}"),("w₂",f"{cub['weights'][1]:.6f}"),
      ("w₃",f"{cub['weights'][2]:.6f}"),("b",f"{cub['bias']:.6f}"),
      ("训练 MSE",f"{cub['mse']:.6f}"),("参数更新次数",cub["steps"])])
    fig("04_cubic_regression.png","图 4  三次多项式回归拟合曲线与训练样本")
    text(f"在同一数据集上，MSE 从 {lin['mse']:.6f} 降至 {cub['mse']:.6f}，降低 {(1-cub['mse']/lin['mse'])*100:.1f}%。"
         "三次曲线能表现弯曲形状，但其结构不是正弦函数本身；噪声不可由数据中完全消除。")
    fig("05_linear_vs_cubic.png","图 5  模型表达能力对比")
    fig("07_regression_training_loss.png","图 6  两种模型优化过程中的 MSE 下降曲线")
    text("核心实现步骤：固定种子→构造 x/y→根据模型生成 (100,1) 或 (100,3) 张量→实例化 Linear→MSELoss→Adam→每轮 zero_grad/backward/step→保存真实参数与预测图。")
else:text("回归部分尚未运行；执行 run_all.py 生成真实结果后再生成报告。")
page()
h("四、MNIST Softmax 分类")
if MN:
    soft=MN["softmax"];hst=soft["history"];last=hst[-1]
    text("MNIST 包含 60,000 张训练图和 10,000 张测试图，每张灰度图为 28×28。输入张量展平为 784 维；基线模型只使用一层 nn.Linear(784,10)，不加入隐藏层或卷积层，严格对应作业 y=softmax(wx+b)。输出为 10 个类别 logits，CrossEntropyLoss 内部处理 log-softmax，因此训练时不手动调用 softmax。")
    tbl(["核心指标","实测结果"],[
        ("模型可训练参数总量",soft["parameter_count"]),
        ("批大小",soft["batch_size"]),
        ("轮次 Epoch",soft["epochs"]),
        ("累计迭代 Iteration",last["iteration"]),
        ("最终训练准确率",f"{last['train_accuracy']*100:.2f}%"),
        ("最终测试准确率",f"{last['test_accuracy']*100:.2f}%"),
        ("最终训练平均交叉熵",f"{last['train_loss']:.5f}"),
        ("最终测试平均交叉熵",f"{last['test_loss']:.5f}"),
        ("优化器",f"{soft['optimizer']} (lr={soft['learning_rate']})")])
    fig("08_softmax_accuracy.png","图 7  随累计训练迭代数变化的训练/测试分类准确率")
    fig("08_softmax_loss.png","图 8  随累计训练迭代数变化的训练/测试平均分类损失")
    text("实现步骤：使用 torchvision 的 MNIST 原始分区；将 uint8 灰度图缩放至 [0,1]；以批为单位 shuffle 并训练；每完成一个 Epoch，在不更新参数的前提下分别评估全量训练集和测试集；将评估点放在累计梯度更新次数轴上，生成两张双曲线图。评估模式下使用 model.eval() 与 torch.inference_mode()，不执行 backward/step。")
    text("分析：准确率表示预测正确样本占比；交叉熵还反映预测概率的置信程度，两者并不必然同步变化。训练集指标与测试集指标存在差距，主要因为训练数据用于优化，而独立测试数据只用于评估泛化能力。")
else:
    text("真实 MNIST 尚未完成，未填写假数据。请运行 run_all.py 并重新执行 create_report.py。")
page()
h("五、拓展实验（非作业必做项目）")
h("5.1 自动求导的数值梯度验证",2)
if r:
    g=r["gradient_check"]
    text("为核验 PyTorch 的反向传播机制，在独立的双精度小样本上比较 autograd(w) 与中心差分 [L(w+h)−L(w−h)]/(2h)。")
    tbl(["项目","结果"],[
        ("Autograd dL/dw",f"{g['autograd_dw']:.10f}"),
        ("中心差分 dL/dw",f"{g['finite_difference_dw']:.10f}"),
        ("相对误差",f"{g['relative_error']:.3e}")])
    text("二者高度一致，说明所设计算图及其求导方向得到了独立数值方法支持。")
h("5.2 模型复杂度与分布外外推",2)
if r:
    ex=r["polynomial_extrapolation"]
    tbl(["多项式阶数","训练 MSE","区间内对 sin(x) MSE","区间外对 sin(x) MSE"],[
        (k,f"{v['training_mse_noisy']:.5f}",f"{v['interpolation_clean_mse']:.5f}",
         f"{v['extrapolation_clean_mse']:.4f}") for k,v in ex.items()])
    fig("06_polynomial_extrapolation.png","图 9  不同次数多项式：区间内拟合与超出 [-3,3] 后的外推")
    text("说明：此处不同阶数通过 NumPy 多项式最小二乘拟合，只用于研究，不替代必做的 PyTorch 三次回归。训练区间外，复杂多项式可能剧烈偏离真实正弦曲线，提示拟合训练数据与可靠泛化是两个不同问题。")
if MN:
    diag=MN["softmax_diagnostics"]
    h("5.3 错误结构、置信度与可解释性",2)
    fig("10_softmax_confusion_matrix.png","图 10  Softmax 在独立测试集上的按真实类别归一化混淆矩阵")
    confs=diag["top_confusions"][:3]
    text("较常见的错分方向（真实→预测）："+"；".join(f"{x['true']}→{x['predicted']}（{x['count']} 个）" for x in confs)+"。该分析帮助定位模型更难区分的笔画组合。")
    fig("11_softmax_high_confidence_errors.png","图 11  模型置信度最高的 20 个错误预测")
    fig("12_softmax_calibration.png","图 12  置信度与实际正确频率之间的一致性")
    text(f"10 等宽区间 Expected Calibration Error（ECE）={diag['ece_10_bins']:.5f}。准确率高并不代表概率可靠：模型对错题仍可能给出很高置信度。ECE 依赖分桶方式，因此作为诊断而不是绝对评分。")
    fig("13_softmax_weight_maps.png","图 13  十个类别的线性权重空间热力图")
    text("每张权重图将 Linear(784,10) 中对应类别的 784 个像素权重重排为 28×28。红色与蓝色分别表示对该类别 logit 的正负影响，这些图不是模型生成的手写数字，也不能直接解释输入像素间的非线性交互。")
    if "cnn_extension" in MN:
        extra=MN["cnn_extension"]
        h("5.4 额外模型：CNN 与线性分类器对比",2)
        tbl(["属性","Softmax 基线","独立 CNN 拓展"],[
          ("测试准确率",f"{soft['final_test']['accuracy']*100:.2f}%",f"{extra['final_test']['accuracy']*100:.2f}%"),
          ("测试平均交叉熵",f"{soft['final_test']['loss']:.5f}",f"{extra['final_test']['loss']:.5f}"),
          ("参数量",str(soft["parameter_count"]),str(extra["parameter_count"])),
          ("训练 Epoch",str(soft["epochs"]),str(extra["epochs"]))])
        fig("15_softmax_vs_cnn_extension.png","图 14  两类模型的测试准确率对比（训练轮次不同，非严格公平基准）")
        text("CNN 利用局部卷积及池化的结构先验来捕获笔画形状。两模型训练轮次和优化器不同，因此对比仅用于启发：不能将所有性能差异归因于网络结构。必做模型依然是 784→10 的线性 Softmax。")
h("六、实验结论与反思")
text("本实验从形状操作开始，建立了数据→特征→预测→损失→梯度→参数更新→独立测试的完整实践链。线性回归揭示表达能力不足的欠拟合，三次特征映射解释非线性特征与线性模型的区别；MNIST 通过 784×10 权重将图像映射到十类别分数，交叉熵指导参数更新。拓展实验进一步提醒：低训练误差不等于外推可靠，高分类准确率也不意味着置信度充分可信。")
h("七、复现方法与文件清单")
code("""cd ~/Desktop/NN_Experiment_1_2026
python run_all.py --device cuda
python create_report.py""")
text("关键文件：numpy_exercise.py（所有 NumPy 必做步骤）；run_all.py（训练、评估、图表与拓展）；results/metrics.json（原始指标）；results/training.log（实际训练记录）；results/softmax_weights.pt（模型权重）；figures/（各图）；README.md（复现与提交说明）。")
text("原始实验要求：https://github.com/caixd-220529/deel_learning_2026-public/blob/main/experiments/exp1/实验一.md")
text("正式提交前：填写点名册序号、班级全称、12 位学号、姓名；核对截止时间；依教师规范改名为 序号-班级-学号-姓名-exp1-v1.docx。")
# Header/footer
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("神经网络实验一  ·  2026  ·  真实运行结果及独立拓展").font.size=Pt(8)
doc.save(DOC)
print("DOCX_CREATED:",DOC, "bytes", DOC.stat().st_size, "inline_shapes",len(doc.inline_shapes))
